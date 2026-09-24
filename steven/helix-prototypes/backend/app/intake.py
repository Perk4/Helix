"""Build a Study Evidence Package from uploaded files.

The workbench seeds one package from a JSON fixture. This module is the other
way in: a person uploads the study's own files and the package is derived from
them, so the records under every later claim are ones somebody can open.

WHAT IS ADMITTED IS NOT WHAT IS READ

Files fall into three outcomes and the receipt says which each one got:

  DATA       CSV or Excel whose name matches a domain below. Parsed into
             `StudyRecords`, and every number a section later computes traces
             back to it.
  AUTHORITY  PDF or Word. Checksummed into the frozen manifest with an
             authority tier and parsed by nothing. A signed protocol governs a
             study without being machine-readable; turning its prose into
             numbers would manufacture evidence.
  REJECTED   an unknown extension, or a tabular file under a name no domain
             claims. `weights.csv` might be body weights and might not. The
             domain comes from the file name, never the contents, because a
             wrong guess mislabels every number derived from it.

WHAT THE FILES CANNOT SAY

Route of administration and protocol version are in no column. They are
required fields on the upload and are recorded with the caller named as their
source. Nothing here infers a governed fact from data.

WHAT THIS DELIBERATELY DOES NOT DO

It produces records, a study, and a frozen manifest. It produces no claims:
computing those is `section_executor`'s job, and a package that arrived with
claims already in it would have no provenance anybody could check. An uploaded
study therefore starts with an empty claim set and a blocked release gate.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from .reporting import load_report_template
from .schemas import (
    Animal,
    DoseGroup,
    GateDecision,
    GateStatus,
    ManifestEntry,
    Measurement,
    MicroscopicFinding,
    ReportSection,
    SectionStatus,
    Study,
    StudyEvidencePackage,
    StudyRecords,
)

SYNTHETIC_LABEL = "SYNTHETIC / NOT FOR SUBMISSION"

# Source tables, by file stem. The suffix is not fixed: a study may arrive as
# CSV or as Excel and both are read into the same rows.
DOMAIN_STEMS = {
    "animal_roster": "animals",
    "body_weights": "body_weights",
    "clinical_observations": "clinical_observations",
    "dosing_formulation": "formulation",
    "food_consumption": "food_consumption",
    "gross_pathology": "gross_pathology",
    "histopathology": "microscopic_findings",
    "organ_weights": "organ_weights",
}

# `StudyRecords` names exactly these. `gross_pathology` has nowhere to go, so
# it is read, reported, and then dropped rather than silently discarded.
PACKAGE_DOMAINS = ("animals", "body_weights", "clinical_observations",
                   "food_consumption", "organ_weights", "microscopic_findings",
                   "formulation")

TABULAR_SUFFIXES = {".csv", ".xlsx", ".xls"}
AUTHORITY_SUFFIXES = {".pdf", ".docx", ".doc"}
ARCHIVE_SUFFIXES = {".zip"}

PROTOCOL_PATTERN = re.compile(r"protocol", re.I)
STATISTICS_PATTERN = re.compile(r"stat|\bsap\b", re.I)

MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_TOTAL_BYTES = 256 * 1024 * 1024
MAX_FILES = 200

NO_FINDING = {"no remarkable findings", "no abnormal findings",
              "no abnormality detected", "no remarkable gross findings",
              "not examined", ""}
CANONICAL_NO_FINDING = "No abnormality detected"

_ID_SAFE = re.compile(r"[^A-Z0-9]+")


class IntakeRejected(ValueError):
    """The upload cannot be accepted, with a reason the caller can act on."""


def _token(text: str) -> str:
    return _ID_SAFE.sub("-", str(text).upper()).strip("-")


def _safe_name(name: str) -> str:
    """Base name only. Directory components in an upload are not honoured."""
    return Path(str(name).replace("\\", "/")).name


def _number(raw, default=None):
    try:
        return float(str(raw).strip())
    except (TypeError, ValueError):
        return default


def _sex(raw: str) -> str:
    value = (raw or "").strip().upper()
    return "M" if value.startswith("M") else "F" if value.startswith("F") else ""


def sha256_of(payload: bytes) -> str:
    return "sha256:" + hashlib.sha256(payload).hexdigest()


# ── classification ────────────────────────────────────────────────────────────
@dataclass
class AcceptedFile:
    name: str
    kind: str                      # data | authority
    domain: str | None
    authority_tier: int
    readable: bool
    bytes: int
    checksum: str

    def as_record(self) -> dict:
        return {"name": self.name, "kind": self.kind, "domain": self.domain,
                "authority_tier": self.authority_tier, "readable": self.readable,
                "bytes": self.bytes, "checksum": self.checksum}


@dataclass
class IntakeReport:
    accepted: list[AcceptedFile] = field(default_factory=list)
    rejected: list[dict] = field(default_factory=list)
    anomalies: list[dict] = field(default_factory=list)

    def note(self, kind: str, detail: str, **extra) -> None:
        self.anomalies.append({"kind": kind, "detail": detail, **extra})

    @property
    def domains(self) -> set[str]:
        return {f.domain for f in self.accepted if f.domain}

    @property
    def missing_required(self) -> list[str]:
        return sorted({"animals", "body_weights"} - self.domains)

    def as_record(self) -> dict:
        return {
            "accepted": [f.as_record() for f in self.accepted],
            "rejected": self.rejected,
            "anomalies": self.anomalies,
            "domains_recognised": sorted(self.domains),
            "missing_required_domains": self.missing_required,
            "data_files": sum(1 for f in self.accepted if f.kind == "data"),
            "authority_files": sum(1 for f in self.accepted if f.kind == "authority"),
        }


def classify(name: str) -> tuple[str, str | None, int, bool, str]:
    """Decide what one file is. Returns (kind, domain, tier, readable, reason)."""
    safe = _safe_name(name)
    suffix = Path(safe).suffix.lower()

    if suffix in AUTHORITY_SUFFIXES:
        tier = 2 if (PROTOCOL_PATTERN.search(safe) or STATISTICS_PATTERN.search(safe)) else 3
        return ("authority", None, tier, False,
                f"{suffix} is an authority document; it is checksummed into the "
                f"manifest and read by nothing")

    if suffix in TABULAR_SUFFIXES:
        stem = re.sub(r"[^a-z0-9]+", "_", Path(safe).stem.lower()).strip("_")
        domain = DOMAIN_STEMS.get(stem)
        if domain is None:
            return ("rejected", None, 0, False,
                    f"'{safe}' is tabular but its name matches no known domain. "
                    f"Expected one of: {', '.join(sorted(DOMAIN_STEMS))}. The "
                    f"domain is not guessed from the contents, because a wrong "
                    f"guess would mislabel every number derived from it.")
        return ("data", domain, 1, True, "")

    return ("rejected", None, 0, False,
            f"'{safe}' has extension '{suffix or '(none)'}', which is not accepted. "
            f"Allowed: {', '.join(sorted(TABULAR_SUFFIXES | AUTHORITY_SUFFIXES | ARCHIVE_SUFFIXES))}")


def expand(name: str, payload: bytes) -> list[tuple[str, bytes]]:
    """One upload becomes one or more (name, bytes). Archives are opened.

    Members are taken by base name, so an entry spelled to climb out of a
    destination cannot; there is nowhere for it to go.
    """
    if Path(_safe_name(name)).suffix.lower() not in ARCHIVE_SUFFIXES:
        return [(_safe_name(name), payload)]
    members: list[tuple[str, bytes]] = []
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            for info in archive.infolist():
                if info.is_dir():
                    continue
                if info.file_size > MAX_FILE_BYTES:
                    raise IntakeRejected(
                        f"{info.filename} is {info.file_size} bytes, over the "
                        f"{MAX_FILE_BYTES} limit")
                members.append((_safe_name(info.filename), archive.read(info)))
    except zipfile.BadZipFile as exc:
        raise IntakeRejected(f"{name} is not a readable zip archive: {exc}") from None
    if not members:
        raise IntakeRejected(f"{name} contains no files")
    return members


# ── reading tables ────────────────────────────────────────────────────────────
def read_table(name: str, payload: bytes) -> list[dict]:
    if Path(name).suffix.lower() == ".csv":
        # utf-8-sig: these exports carry a BOM, which corrupts the first header.
        text = payload.decode("utf-8-sig", errors="replace")
        return [dict(row) for row in csv.DictReader(io.StringIO(text))]
    return _read_excel(payload)


def _read_excel(payload: bytes) -> list[dict]:
    """First worksheet, first row as the header.

    Values are rendered back to strings so an Excel study and a CSV study reach
    the parsers below identically. Two typed paths would mean two rounding
    behaviours, and the file format must not be able to change the evidence.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
        except StopIteration:
            return []
        out = []
        for row in rows:
            if all(cell is None for cell in row):
                continue
            record = {}
            for key, cell in zip(header, row):
                if not key:
                    continue
                if cell is None:
                    record[key] = ""
                elif isinstance(cell, float) and cell.is_integer():
                    record[key] = str(int(cell))
                else:
                    record[key] = str(cell).strip()
            out.append(record)
        return out
    finally:
        workbook.close()


# ── domain parsers ────────────────────────────────────────────────────────────
def _animals(rows: list[dict], study_id: str, report: IntakeReport) -> list[Animal]:
    animals = []
    for index, row in enumerate(rows, start=1):
        sex = _sex(row.get("sex"))
        if not sex:
            report.note("unreadable_sex",
                        f"animal {row.get('animal_id')} has sex {row.get('sex')!r}; "
                        f"the row is dropped because every record is joined on sex")
            continue
        animals.append(Animal(
            animal_id=row["animal_id"], study_id=study_id,
            group_id="G" + str(row["group_number"]).strip(), sex=sex,
            randomization_id=f"RAND-{row['group_number']}-{sex}-{index}"))
    return animals


def _body_weights(rows: list[dict], report: IntakeReport) -> list[Measurement]:
    out = []
    for row in rows:
        day, value = _number(row.get("study_day")), _number(row.get("body_weight_g"))
        if day is None or value is None:
            report.note("incomplete_body_weight",
                        f"animal {row.get('animal_id')} day {row.get('study_day')!r}: "
                        f"dropped, because a weight with no day cannot be placed "
                        f"on the study timeline")
            continue
        day, animal = int(day), row["animal_id"]
        out.append(Measurement(
            record_id=f"BW-{animal}-{day}", domain="BW", animal_id=animal,
            timepoint=f"DAY {day}", test_code="BW", value=value, unit="g",
            grain="animal_x_day", source_pointer=f"A-BW#{animal}:DAY{day}"))
    return out


def _organ_weights(rows: list[dict], report: IntakeReport) -> list[Measurement]:
    out = []
    for row in rows:
        value = _number(row.get("weight_g"))
        animal, organ = row["animal_id"], _token(row.get("organ"))
        if value is None:
            report.note("missing_organ_weight",
                        f"animal {animal} organ {row.get('organ')}: no weight recorded")
            continue
        out.append(Measurement(
            record_id=f"OM-{animal}-{organ}", domain="OM", animal_id=animal,
            timepoint="TERMINAL", test_code=organ, value=value, unit="g",
            grain="animal", source_pointer=f"A-OM#{animal}:{organ}"))
    return out


def _clinical_observations(rows: list[dict], report: IntakeReport) -> list[Measurement]:
    """One record per observation period, which is the grain the file records."""
    out = []
    for index, row in enumerate(rows, start=1):
        animal = row["animal_id"]
        start = _number(row.get("study_day_start"))
        if start is None:
            report.note("observation_without_interval",
                        f"animal {animal}: observation has no study_day_start")
            continue
        start = int(start)
        end = int(_number(row.get("study_day_end"), start))
        label = (row.get("period_label") or f"DAYS {start}-{end}").strip()
        finding = (row.get("finding") or "").strip()
        # Studies spell a normal observation differently. Normalising the
        # wording is safe; leaving it alone made every normal row read as a
        # clinical sign the first time this ran.
        value = finding if finding.lower() not in NO_FINDING else CANONICAL_NO_FINDING
        out.append(Measurement(
            record_id=f"CL-{animal}-{start}-{end}-{index}", domain="CL",
            animal_id=animal, timepoint=label.upper(), test_code="CLOBS",
            value=value, unit=None, grain="animal_x_interval",
            source_pointer=f"A-CL#{animal}:{_token(label)}"))
    return out


def grade_scale(rows: list[dict]) -> dict[str, set[str]]:
    """The grade-to-word mapping this study actually uses.

    Derived per study, because studies do not share a scale: one reads
    1=Minimal 2=Mild 3=Moderate and another reads 1=Mild 2=Moderate 3=Marked.
    Both are internally consistent, and judging either against a fixed ranking
    produces nothing but false positives.
    """
    scale: dict[str, set[str]] = {}
    for row in rows:
        grade = (row.get("grade") or "").strip()
        text = (row.get("grade_text") or "").strip().lower()
        if grade and text:
            scale.setdefault(grade, set()).add(text)
    return scale


def _microscopic(rows: list[dict], report: IntakeReport) -> list[MicroscopicFinding]:
    scale = grade_scale(rows)
    for grade, texts in sorted(scale.items()):
        if len(texts) > 1:
            report.note("inconsistent_grade_scale",
                        f"grade {grade} is labelled {sorted(texts)} in different "
                        f"rows, so no severity can be read from the grade alone",
                        grade=grade)

    out, seen = [], {}
    for row in rows:
        animal, tissue = row["animal_id"], (row.get("tissue") or "").strip()
        finding = (row.get("finding") or "").strip()
        is_finding = finding.lower() not in NO_FINDING
        text = (row.get("grade_text") or "").strip().lower()
        if not is_finding:
            severity = "none"
        elif text:
            severity = text
        else:
            severity = "ungraded"
            report.note("graded_without_text",
                        f"{animal}/{tissue}/{finding}: grade {row.get('grade')!r} "
                        f"carries no grade_text, so its severity cannot be named")

        base = f"MI-{animal}-{_token(tissue)}"
        seen[base] = seen.get(base, 0) + 1
        out.append(MicroscopicFinding(
            finding_id=base if seen[base] == 1 else f"{base}-{seen[base]}",
            domain="MI", animal_id=animal, tissue=tissue,
            finding=finding if is_finding else CANONICAL_NO_FINDING,
            severity=severity,
            # The pathologist's own term, upper-cased. Not slugged: rewriting it
            # would quietly invent vocabulary.
            controlled_term=finding.upper() if is_finding else "NORMAL",
            source_pointer=f"A-MI#{animal}:{tissue}"))
    return out


def _formulation(rows: list[dict], report: IntakeReport) -> list[Measurement]:
    """Analytical concentration per group per sampling day.

    'ND' is a result, not a gap. In a vehicle control, not detecting the test
    article is the expected outcome; recording it as missing data would confuse
    absent with not-required. It is only worth reporting in a dosed group.
    """
    out = []
    for row in rows:
        raw = (row.get("measured_conc_mgml") or "").strip()
        group = "G" + str(row["group_number"]).strip()
        day, dose = _number(row.get("study_day")), _number(row.get("dose_mgkg_day"), 0.0)
        value = _number(raw)
        below_detection = value is None and raw.upper() in {"ND", "N/A", "BLQ"}
        if day is None or (value is None and not below_detection):
            report.note("incomplete_formulation",
                        f"{group} day {row.get('study_day')!r}: measured "
                        f"concentration {raw!r} cannot be read as a result")
            continue
        if below_detection and dose:
            report.note("test_article_not_detected_in_dosed_group",
                        f"{group} (dose {dose}) day {int(day)}: reported as {raw!r}")
        out.append(Measurement(
            record_id=f"PC-{group}-D{int(day)}", domain="PC", group_id=group,
            timepoint=f"DAY {int(day)}", test_code="CONC",
            # Below detection keeps the lab's own token. A zero would assert a
            # measured concentration of zero.
            value=raw if below_detection else value,
            unit="mg/mL", grain="concentration_x_timepoint",
            source_pointer=f"A-PC#{group}:DAY{int(day)}"))
    return out


def _food_consumption(rows: list[dict], _report: IntakeReport) -> list[Measurement]:
    out = []
    for index, row in enumerate(rows, start=1):
        value = _number(row.get("food_g_per_animal_per_day") or row.get("value"))
        if value is None:
            continue
        group = "G" + str(row.get("group_number", "")).strip()
        week = str(row.get("study_week") or row.get("week") or index).strip()
        out.append(Measurement(
            record_id=f"FW-{group}-W{week}", domain="FW", group_id=group,
            timepoint=f"WEEK {week}", test_code="FOOD", value=value,
            unit="g/animal/day", grain="group_x_week",
            source_pointer=f"A-FW#{group}:WEEK{week}"))
    return out


PARSERS = {
    "body_weights": _body_weights,
    "organ_weights": _organ_weights,
    "clinical_observations": _clinical_observations,
    "microscopic_findings": _microscopic,
    "formulation": _formulation,
    "food_consumption": _food_consumption,
}


# ── the package ───────────────────────────────────────────────────────────────
def _dose_groups(roster: list[dict], study_id: str,
                 report: IntakeReport) -> list[DoseGroup]:
    seen: dict[str, dict] = {}
    for row in roster:
        group_id = "G" + str(row["group_number"]).strip()
        entry = seen.setdefault(group_id, {
            "label": (row.get("group_name") or group_id).strip(),
            "dose": _number(row.get("dose_mgkg_day"), 0.0),
            "sexes": set(), "counts": {}})
        sex = _sex(row.get("sex"))
        if sex:
            entry["sexes"].add(sex)
            entry["counts"][sex] = entry["counts"].get(sex, 0) + 1

    groups = []
    for group_id in sorted(seen, key=lambda g: int(g[1:])):
        entry = seen[group_id]
        if len(set(entry["counts"].values())) > 1:
            report.note("unequal_group_sizes",
                        f"{group_id} is not balanced by sex: {entry['counts']}. "
                        f"planned_n_per_sex reports the smallest, so a coverage "
                        f"rule cannot silently pass on the larger sex.")
        groups.append(DoseGroup(
            group_id=group_id, label=entry["label"], dose=entry["dose"],
            study_id=study_id, dose_unit="mg/kg/day",
            sexes=sorted(entry["sexes"]),
            planned_n_per_sex=min(entry["counts"].values())))
    return groups


def build_package(
    *, study_id: str, uploads: list[tuple[str, bytes]], route: str,
    protocol_version: str, authorized_by: str, study_type_id: str,
    study_start: str = "",
) -> tuple[StudyEvidencePackage, IntakeReport]:
    """Turn uploaded files into a Study Evidence Package with no claims in it."""
    expanded: list[tuple[str, bytes]] = []
    for name, payload in uploads:
        expanded.extend(expand(name, payload))
    if not expanded:
        raise IntakeRejected("no files were uploaded")
    if len(expanded) > MAX_FILES:
        raise IntakeRejected(f"{len(expanded)} files, over the {MAX_FILES} limit")
    total = sum(len(p) for _n, p in expanded)
    if total > MAX_TOTAL_BYTES:
        raise IntakeRejected(f"{total} bytes, over the {MAX_TOTAL_BYTES} limit")

    report = IntakeReport()
    tables: dict[str, list[dict]] = {}
    manifest: list[ManifestEntry] = []
    seen_names: set[str] = set()

    for name, payload in expanded:
        if len(payload) > MAX_FILE_BYTES:
            report.rejected.append({"name": name, "reason": "over the size limit"})
            continue
        if name.lower() in seen_names:
            report.rejected.append({"name": name, "reason": "duplicate file name"})
            continue
        seen_names.add(name.lower())

        kind, domain, tier, readable, reason = classify(name)
        if kind == "rejected":
            report.rejected.append({"name": name, "reason": reason})
            continue

        checksum = sha256_of(payload)
        report.accepted.append(AcceptedFile(
            name=name, kind=kind, domain=domain, authority_tier=tier,
            readable=readable, bytes=len(payload), checksum=checksum))
        manifest.append(ManifestEntry(
            artifact_id=f"A-{_token(Path(name).stem)}",
            kind="source_data" if kind == "data" else
                 ("protocol" if tier == 2 else "reference"),
            name=name,
            version=protocol_version if tier == 2 else "uploaded",
            authority_tier=tier, checksum=checksum, locked=True,
            authorized_by=authorized_by))
        if kind == "data":
            tables[domain] = read_table(name, payload)

    if report.missing_required:
        raise IntakeRejected(
            f"the upload has no {', '.join(report.missing_required)}. A study "
            f"needs a roster to join on and at least one measurement table.")

    animals = _animals(tables.get("animals", []), study_id, report)
    if not animals:
        raise IntakeRejected("the roster produced no usable animals")

    records = StudyRecords(
        animals=animals,
        **{domain: PARSERS[domain](tables.get(domain, []), report)
           for domain in PACKAGE_DOMAINS if domain != "animals"})

    if tables.get("gross_pathology"):
        report.note("domain_dropped",
                    f"gross_pathology has {len(tables['gross_pathology'])} rows and "
                    f"StudyRecords does not name it, so it is not carried into the "
                    f"package. It is listed in the manifest and can still be read.")

    days = sorted({int(m.timepoint.removeprefix("DAY ")) for m in records.body_weights})
    roster = tables.get("animals", [])
    strain = next((r.get("strain", "").strip() for r in roster if r.get("strain")), "")
    species = next((r.get("species", "").strip() for r in roster if r.get("species")), "rat")

    template = load_report_template()
    package = StudyEvidencePackage(
        package_id=f"PKG-{study_id.removeprefix('STUDY-')}",
        label=SYNTHETIC_LABEL,
        workflow_state="gated",
        manifest=manifest,
        study=Study(
            study_id=study_id, study_type_id=study_type_id,
            species=f"{strain} {species}".strip(), route=route,
            duration_days=max(days) if days else 1,
            study_start=study_start or datetime.now(UTC).date().isoformat(),
            protocol_version=protocol_version,
            dose_groups=_dose_groups(roster, study_id, report)),
        records=records,
        report_sections=[ReportSection(
            section_id=section.section_id, template_id=template.template_id,
            title=section.title,
            required_fields=sum(1 for f in section.fields if f.required),
            status=SectionStatus.NEEDS_REVIEW) for section in template.sections],
        # No claims. Computing them is the section executor's job, and a package
        # that arrived with claims would have provenance nobody could check.
        claims=[], provenance_edges=[], validation_results=[],
        review_dispositions=[], approvals=[],
        gate_decisions=[GateDecision(
            gate_id="GATE-RELEASE", gate_type="release",
            status=GateStatus.BLOCKED,
            blocking_result_ids=[],
            decided_at=datetime.now(UTC).isoformat())],
        export_artifacts=[], retrieval_index=[], events=[])
    return package, report
